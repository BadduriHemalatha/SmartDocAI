"""Extract text and metadata from PDF, DOCX, and TXT uploads."""

from __future__ import annotations

import io
from pathlib import Path

from smartdoc.exceptions import (
    CorruptedFileError,
    EmptyDocumentError,
    ExtractionError,
    FileTooLargeError,
    UnsupportedFileTypeError,
)
from smartdoc.logging_setup import get_logger
from smartdoc.models import DocumentRecord, ExtractedPage

logger = get_logger()

SUPPORTED = {".pdf", ".docx", ".txt"}


class UploadedFileLike:
    """Minimal protocol implemented by Streamlit UploadedFile and test doubles."""

    name: str
    size: int

    def getvalue(self) -> bytes:  # pragma: no cover - protocol
        raise NotImplementedError


def _suffix(name: str) -> str:
    return Path(name).suffix.lower()


def extract_documents(
    files: list,
    *,
    max_file_size_bytes: int,
) -> list[DocumentRecord]:
    if not files:
        raise ExtractionError(
            "No files provided",
            user_message="Please upload at least one PDF, DOCX, or TXT document.",
        )

    records: list[DocumentRecord] = []
    for uploaded in files:
        records.append(extract_one(uploaded, max_file_size_bytes=max_file_size_bytes))
    return records


def extract_one(uploaded, *, max_file_size_bytes: int) -> DocumentRecord:
    name = getattr(uploaded, "name", None) or "unnamed"
    suffix = _suffix(name)

    if suffix not in SUPPORTED:
        raise UnsupportedFileTypeError(
            f"Unsupported type: {suffix}",
            user_message=(
                f"'{name}' is not a supported file type. "
                "Please upload PDF, DOCX, or TXT files only."
            ),
        )

    data = _read_bytes(uploaded)
    size = len(data)

    if size > max_file_size_bytes:
        mb = max_file_size_bytes / (1024 * 1024)
        raise FileTooLargeError(
            f"File too large: {name} ({size} bytes)",
            user_message=f"'{name}' exceeds the maximum size of {mb:.0f} MB.",
        )

    if size == 0:
        raise EmptyDocumentError(
            f"Empty file: {name}",
            user_message=f"'{name}' is empty and cannot be processed.",
        )

    try:
        if suffix == ".pdf":
            pages = _extract_pdf(name, data)
        elif suffix == ".docx":
            pages = _extract_docx(name, data)
        else:
            pages = _extract_txt(name, data)

    except (UnsupportedFileTypeError, EmptyDocumentError, CorruptedFileError):
        raise

    except Exception as exc:
        logger.exception("Extraction failed for %s", name)
        raise ExtractionError(
            f"Failed to extract {name}: {exc}",
            user_message=(
                f"Could not read '{name}'. The file may be corrupted or password-protected."
            ),
        ) from exc

    combined = "\n".join(
        p.text.strip()
        for p in pages
        if p.text and p.text.strip()
    )

    if not combined.strip():
        raise EmptyDocumentError(
            f"No extractable text in {name}",
            user_message=(
                f"No readable text was found in '{name}'. "
                "The PDF may contain scanned images that could not be read by OCR."
            ),
        )

    logger.info(
        "Extracted %s: %s pages/units, %s characters",
        name,
        len(pages),
        len(combined),
    )

    return DocumentRecord(
        document_name=name,
        file_type=suffix.lstrip("."),
        size_bytes=size,
        page_count=len(pages),
        char_count=len(combined),
        pages=pages,
    )


def _read_bytes(uploaded) -> bytes:
    if hasattr(uploaded, "getvalue"):
        return uploaded.getvalue()

    if hasattr(uploaded, "read"):
        data = uploaded.read()
        if isinstance(data, str):
            return data.encode("utf-8")
        return data

    raise ExtractionError("Upload object has no readable content.")


def _extract_pdf(name: str, data: bytes) -> list[ExtractedPage]:
    """Extract normal PDF text and use Tesseract OCR as a fallback for scanned pages."""

    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("pypdf is not installed.") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except PdfReadError as exc:
        raise CorruptedFileError(
            f"Corrupt PDF: {name}",
            user_message=(
                f"'{name}' could not be opened as a PDF. The file may be corrupted."
            ),
        ) from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:
            raise CorruptedFileError(
                f"Encrypted PDF: {name}",
                user_message=(
                    f"'{name}' is password-protected and cannot be processed."
                ),
            ) from exc

    pages: list[ExtractedPage] = []

    # OCR libraries are imported only for PDF OCR fallback.
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
    except ImportError:
        pymupdf = None
        pytesseract = None
        Image = None

    # Use the known Tesseract installation on Windows.
    if pytesseract is not None:
        tesseract_path = Path(
            r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        )

        if tesseract_path.exists():
            pytesseract.pytesseract.tesseract_cmd = str(tesseract_path)

    # Open the PDF once with PyMuPDF if OCR support is available.
    ocr_document = None

    if pymupdf is not None and pytesseract is not None and Image is not None:
        try:
            ocr_document = pymupdf.open(
                stream=data,
                filetype="pdf",
            )
        except Exception as exc:
            logger.warning(
                "Could not open %s for OCR: %s",
                name,
                exc,
            )

    for index, page in enumerate(reader.pages, start=1):

        # First try the existing normal PDF text extraction.
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            logger.warning(
                "Failed extracting page %s of %s: %s",
                index,
                name,
                exc,
            )
            text = ""

        # If normal extraction finds no text, use OCR.
        if not text.strip() and ocr_document is not None:
            try:
                pdf_page = ocr_document[index - 1]

                # Render page at 2x resolution for better OCR accuracy.
                pixmap = pdf_page.get_pixmap(
                    matrix=pymupdf.Matrix(2, 2),
                    alpha=False,
                )

                image_bytes = pixmap.tobytes("png")

                with Image.open(io.BytesIO(image_bytes)) as image:
                    ocr_text = pytesseract.image_to_string(
                        image,
                        lang="eng",
                    )

                text = ocr_text or ""

                if text.strip():
                    logger.info(
                        "OCR extracted text from page %s of %s",
                        index,
                        name,
                    )
                else:
                    logger.warning(
                        "OCR found no text on page %s of %s",
                        index,
                        name,
                    )

            except Exception as exc:
                logger.warning(
                    "OCR failed for page %s of %s: %s",
                    index,
                    name,
                    exc,
                )

        pages.append(
            ExtractedPage(
                document_name=name,
                file_type="pdf",
                text=text,
                page_number=index,
            )
        )

    if ocr_document is not None:
        try:
            ocr_document.close()
        except Exception:
            pass

    if not pages:
        raise EmptyDocumentError(
            f"PDF has no pages: {name}",
            user_message=f"'{name}' does not contain any pages.",
        )

    return pages


def _extract_docx(name: str, data: bytes) -> list[ExtractedPage]:
    try:
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("python-docx is not installed.") from exc

    try:
        document = Document(io.BytesIO(data))
    except PackageNotFoundError as exc:
        raise CorruptedFileError(
            f"Corrupt DOCX: {name}",
            user_message=f"'{name}' could not be opened as a Word document.",
        ) from exc
    except Exception as exc:
        raise CorruptedFileError(
            f"Corrupt DOCX: {name}",
            user_message=f"'{name}' could not be opened as a Word document.",
        ) from exc

    paragraphs = [
        p.text.strip()
        for p in document.paragraphs
        if p.text and p.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text and cell.text.strip()
            ]

            if cells:
                paragraphs.append(" | ".join(cells))

    text = "\n".join(paragraphs)

    return [
        ExtractedPage(
            document_name=name,
            file_type="docx",
            text=text,
            page_number=None,
        )
    ]


def _extract_txt(name: str, data: bytes) -> list[ExtractedPage]:
    text = _decode_text(data)
    lines = text.splitlines()

    return [
        ExtractedPage(
            document_name=name,
            file_type="txt",
            text=text,
            page_number=None,
            start_line=1 if lines else None,
            end_line=len(lines) if lines else None,
        )
    ]


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    return data.decode("utf-8", errors="replace")